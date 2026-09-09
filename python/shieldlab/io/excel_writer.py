from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from shieldlab.analysis import compare_coefficients, hvl, load_reference_coefficients, mean_free_path, tvl
from shieldlab.core.descriptors import descriptors_from_study, elemental_expansion_table
from shieldlab.plotting import create_standard_plots


def _read_optional_csv(path: Path) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame()
    return pd.read_csv(path)


def _study_info(study_file: str | Path | None) -> pd.DataFrame:
    if study_file is None:
        return pd.DataFrame()
    study_path = Path(study_file)
    if not study_path.exists():
        return pd.DataFrame()
    with study_path.open("r", encoding="utf-8-sig") as handle:
        study = json.load(handle)

    rows: list[dict[str, object]] = []
    rows.append({"field": "name", "value": study.get("name", "")})
    rows.append({"field": "description", "value": study.get("description", "")})
    for section_name in ("source", "geometry", "run"):
        section = study.get(section_name, {})
        rows.append({"field": section_name, "value": json.dumps(section, ensure_ascii=False)})
    return pd.DataFrame(rows)


def _derived_summary(summary: pd.DataFrame, layers: pd.DataFrame) -> pd.DataFrame:
    if summary.empty:
        return pd.DataFrame()
    row = summary.iloc[0].to_dict()
    mu = float(row.get("linear_attenuation_cm_inv", 0.0))
    density = None
    if len(layers.index) == 1:
        density = float(layers.iloc[0].get("density_g_cm3", 0.0))
    mass_attenuation = mu / density if density and density > 0 else None

    return pd.DataFrame(
        [
            {
                "linear_attenuation_cm_inv": mu,
                "mass_attenuation_cm2_g": mass_attenuation,
                "mean_free_path_cm": mean_free_path(mu),
                "hvl_cm": hvl(mu),
                "tvl_cm": tvl(mu),
                "attenuation_estimate_type": row.get("attenuation_estimate_type", "direct"),
                "transmission_fraction": row.get("transmission_fraction", 0.0),
                "reflection_fraction": row.get("reflection_fraction", 0.0),
                "absorption_fraction": row.get("absorption_fraction", 0.0),
            }
        ]
    )


def _figure_index(figures: list[Path], result_path: Path) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "figure": figure.name,
                "relative_path": figure.relative_to(result_path).as_posix(),
            }
            for figure in figures
        ]
    )


# ── Styling helpers ─────────────────────────────────────────────────────────

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_ALT_FILL = PatternFill("solid", fgColor="D9E1F2")
_WARN_FILL = PatternFill("solid", fgColor="FFEB9C")   # abs(pct diff) > 5 %
_ERR_FILL  = PatternFill("solid", fgColor="FFC7CE")   # abs(pct diff) > 15 %
_CENTERED  = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _style_sheet(
    ws: Any,
    freeze_row: int = 1,
    pct_diff_column: str | None = None,
) -> None:
    """Apply header formatting, alternate row shading, frozen pane, and column widths."""
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _CENTERED

    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        fill = _ALT_FILL if row_idx % 2 == 0 else PatternFill()
        for cell in row:
            if cell.fill == PatternFill():
                cell.fill = fill
            cell.alignment = Alignment(horizontal="left", vertical="center")

    if pct_diff_column:
        col_names = [cell.value for cell in ws[1]]
        if pct_diff_column in col_names:
            col_idx = col_names.index(pct_diff_column) + 1
            for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                for cell in row:
                    try:
                        val = float(cell.value)
                        if abs(val) > 15:
                            cell.fill = _ERR_FILL
                        elif abs(val) > 5:
                            cell.fill = _WARN_FILL
                    except (TypeError, ValueError):
                        pass

    ws.freeze_panes = f"A{freeze_row + 1}"

    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len + 2, 12), 50)


# ── Workbook builder ─────────────────────────────────────────────────────────


def write_workbook(
    result_dir: str | Path,
    output_file: str | Path | None = None,
    study_file: str | Path | None = None,
    make_plots: bool = True,
) -> Path:
    result_path = Path(result_dir)
    if output_file is None:
        output_path = result_path / "shieldlab_results.xlsx"
    else:
        output_path = Path(output_file)

    summary = _read_optional_csv(result_path / "run_summary.csv")
    layers = _read_optional_csv(result_path / "layer_energy_deposition.csv")
    sweep = _read_optional_csv(result_path / "sweep_summary.csv")
    thickness_sweep = _read_optional_csv(result_path / "thickness_sweep_summary.csv")
    derived = _derived_summary(summary, layers)
    study_info_df = _study_info(study_file)
    reference = load_reference_coefficients(study_file)
    comparison = compare_coefficients(sweep, reference)
    if not reference.empty:
        reference.to_csv(result_path / "reference_coefficients.csv", index=False)
    if not comparison.empty:
        comparison.to_csv(result_path / "reference_comparison.csv", index=False)
    figures = create_standard_plots(result_path) if make_plots else []
    figure_index = _figure_index(figures, result_path)

    # Material descriptors from study JSON
    study_dict: dict[str, Any] = {}
    if study_file is not None:
        sp = Path(study_file)
        if sp.exists():
            with sp.open("r", encoding="utf-8-sig") as fh:
                study_dict = json.load(fh)
    mat_descriptors = descriptors_from_study(study_dict)
    elem_expansion = elemental_expansion_table(study_dict)

    sheet_map: dict[str, tuple[pd.DataFrame, str | None]] = {
        "Study_Info":             (study_info_df,   None),
        "Run_Summary":            (summary,         None),
        "Layer_Edep":             (layers,          None),
        "Derived_Coefficients":   (derived,         None),
        "Material_Descriptors":   (mat_descriptors, None),
        "Elemental_Expansion":    (elem_expansion,  None),
        "Sweep_Summary":          (sweep,           None),
        "Thickness_Sweep":        (thickness_sweep, None),
        "Composition_Sweep":      (_read_optional_csv(result_path / "composition_sweep_summary.csv"), None),
        "Reference_Coefficients": (reference,       None),
        "Reference_Comparison":   (comparison,      "mass_attenuation_cm2_g_percent_difference"),
        "Figures_Index":          (figure_index,    None),
    }

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, (df, pct_col) in sheet_map.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            _style_sheet(writer.sheets[sheet_name], pct_diff_column=pct_col)

    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(description="Create a ShieldLab-G4 Excel workbook from CSV results.")
    parser.add_argument("result_dir", help="Directory containing run_summary.csv and layer_energy_deposition.csv")
    parser.add_argument("--output", help="Output .xlsx path")
    parser.add_argument("--study", help="Optional study JSON file to include as Study_Info")
    parser.add_argument("--no-plots", action="store_true", help="Skip generation of standard PNG figures")
    args = parser.parse_args()
    workbook = write_workbook(args.result_dir, args.output, args.study, make_plots=not args.no_plots)
    print(workbook)


if __name__ == "__main__":
    main()